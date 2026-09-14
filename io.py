"""UTF-8 BOM付きCSVと、納品Excelの入力表を読む。Excel数式は入力に使わない。"""
import csv
import json
from datetime import date, datetime
from pathlib import Path

TABLES = ["staff", "schedule", "activities", "requests", "overtime_candidates", "overtime_ledger", "leave_ledger"]
SHEETS = dict(staff="職員マスタ", schedule="勤務入力", activities="休憩と担当", requests="希望休",
              overtime_candidates="残業候補", overtime_ledger="残業集計", leave_ledger="有給集計")


def load_folder(folder):
    folder = Path(folder)
    data = {"config":json.loads((folder/"config.json").read_text(encoding="utf-8"))}
    for name in TABLES:
        with (folder/f"{name}.csv").open(encoding="utf-8-sig", newline="") as stream:
            data[name] = list(csv.DictReader(stream))
    return data


def write_csv(path, rows, fields):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            # Excelに読み込む文字列を数式として実行させない。
            clean = {k:("'"+v if isinstance(v,str) and v.startswith(("=","+","-","@")) else v) for k,v in row.items()}
            writer.writerow(clean)


def load_excel(path):
    """読み取り専用。出力結果のセル・保存済み判定は信用せず入力表から再監査する。"""
    import openpyxl
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    settings = workbook["施設設定"]
    config = {}
    for key, value, *_ in settings.iter_rows(min_row=6, values_only=True):
        if key:
            if key=='floors':
                config[key]=str(value).split('・')
            elif key=='special_clause':
                config[key]=str(value).lower()=='true'
            else:
                config[str(key)] = value
    config['shifts']={}
    for code,begin,end,next_day,pause,night,*_ in workbook['勤務区分'].iter_rows(min_row=6,values_only=True):
        if code:
            config['shifts'][str(code)]={'start':str(begin)[:5],'end':str(end)[:5],
                'next_day':int(next_day),'break_minutes':int(pause),'night':str(night).lower()=='true'}
    data = {"config":config}
    for key, name in SHEETS.items():
        sheet = workbook[name]
        headers = [c.value for c in next(sheet.iter_rows(min_row=5,max_row=5))]
        rows = []
        for cells in sheet.iter_rows(min_row=6):
            if all(c.value is None for c in cells):
                continue
            row = {}
            for header, cell in zip(headers, cells):
                if not header:
                    continue
                if cell.data_type == "f":
                    raise ValueError(f"{name}!{cell.coordinate}: 入力欄は値で指定してください")
                value = cell.value
                if isinstance(value, (datetime,date)):
                    value = value.date().isoformat() if header in {"date","shift_date","grant_date"} and isinstance(value,datetime) else value.isoformat()
                row[header] = "" if value is None else str(value)
            rows.append(row)
        data[key] = rows
    workbook.close()
    return data
